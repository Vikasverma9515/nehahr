"""CSV import and resume parsing."""

from app.services import intake

NAUKRI_STYLE = (
    "Candidate Name,Mobile No.,Email ID,Current Location,Total Experience,Notice Period,Key Skills\n"
    "Priya Sharma,98765 43210,priya@x.com,Pune,5.5 Years,2 months,\"Python, SQL; AWS\"\n"
    "Rahul Verma,+91-91234-56789,,Bangalore,3,30 days,Java\n"
    "No Phone,,np@x.com,Delhi,,,\n"
    ",,,,,,\n"
).encode()


def test_parse_csv_maps_headers_and_normalizes():
    rows, errors, mapping = intake.parse_csv(NAUKRI_STYLE)
    assert mapping["Candidate Name"] == "name" and mapping["Mobile No."] == "phone"
    assert [r["phone"] for r in rows] == ["+919876543210", "+919123456789"]
    assert rows[0]["experience_years"] == 5.5
    assert rows[0]["notice_period_days"] == 60 and rows[1]["notice_period_days"] == 30
    assert rows[0]["skills"] == ["Python", "SQL", "AWS"]
    assert errors == [{"line": 4, "name": "No Phone", "error": "missing phone"}]


def test_parse_csv_semicolon_and_bom():
    data = "﻿Name;Phone\nAsha;9876500000\n".encode("utf-8")
    rows, errors, _ = intake.parse_csv(data)
    assert rows == [{"name": "Asha", "phone": "+919876500000"}] and not errors


def test_extract_text_docx_and_txt():
    import io
    import docx

    doc = docx.Document()
    doc.add_paragraph("Priya Sharma")
    doc.add_paragraph("Senior Backend Engineer at Acme")
    buf = io.BytesIO()
    doc.save(buf)
    assert "Senior Backend Engineer" in intake.extract_text("cv.docx", buf.getvalue())
    assert intake.extract_text("cv.txt", b"hello   world") == "hello world"


def test_candidate_fields_from_resume():
    fields = intake.candidate_fields_from_resume({
        "name": "Priya", "phone": "98765 43210", "experience_years": 6, "skills": ["Go", "K8s"], "summary": "x",
    })
    assert fields == {"name": "Priya", "phone": "+919876543210", "experience_years": 6, "skills": ["Go", "K8s"]}


def test_match_prompt_formats():
    text = intake.MATCH_PROMPT.format(title="t", role_type="r", work_model="w", skills="s", description="d", resume="{}")
    assert '"score": int' in text


def test_import_endpoint_dedupes(monkeypatch):
    from fastapi.testclient import TestClient
    from app.config import settings
    from app.main import app
    import app.routers.candidates as cand

    monkeypatch.setattr(settings, "disable_auth", True)
    monkeypatch.setattr(cand, "_existing_keys", lambda: ({"+919876543210"}, set()))
    r = TestClient(app).post(
        "/api/candidates/import",
        files={"file": ("c.csv", NAUKRI_STYLE, "text/csv")},
        data={"dry_run": "true"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["would_create"] == 1 and body["duplicates"][0]["name"] == "Priya Sharma"
    assert body["created"] == 0
